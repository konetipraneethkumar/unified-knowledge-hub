import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from app.services.document_extraction import extract_document_text


class DocumentExtractionTests(unittest.TestCase):
	def test_extracts_plain_text_markdown_csv_and_json(self):
		with tempfile.TemporaryDirectory() as temporary_directory:
			root = Path(temporary_directory)
			(root / "note.txt").write_text("plain text", encoding="utf-8")
			(root / "note.md").write_text("# heading", encoding="utf-8")
			(root / "table.csv").write_text("name,value\nalpha,2", encoding="utf-8")
			(root / "record.json").write_text(
				json.dumps({"name": "alpha"}), encoding="utf-8"
			)

			self.assertEqual(extract_document_text(root / "note.txt"), "plain text")
			self.assertEqual(extract_document_text(root / "note.md"), "# heading")
			self.assertEqual(
				extract_document_text(root / "table.csv"), "name\tvalue\nalpha\t2"
			)
			self.assertIn('"alpha"', extract_document_text(root / "record.json"))

	def test_extracts_docx_pptx_and_xlsx_text(self):
		from docx import Document
		from openpyxl import Workbook
		from pptx import Presentation
		from pptx.util import Inches

		with tempfile.TemporaryDirectory() as temporary_directory:
			root = Path(temporary_directory)
			document_path = root / "document.docx"
			document = Document()
			document.add_paragraph("docx marker")
			document.save(document_path)

			presentation_path = root / "slides.pptx"
			presentation = Presentation()
			slide = presentation.slides.add_slide(presentation.slide_layouts[5])
			slide.shapes.add_textbox(
				Inches(1), Inches(1), Inches(4), Inches(1)
			).text = "pptx marker"
			table = slide.shapes.add_table(
				1, 1, Inches(1), Inches(2), Inches(3), Inches(1)
			).table
			table.cell(0, 0).text = "pptx table marker"
			presentation.save(presentation_path)

			workbook_path = root / "sheet.xlsx"
			workbook = Workbook()
			workbook.active["A1"] = "xlsx marker"
			workbook.save(workbook_path)

			self.assertIn("docx marker", extract_document_text(document_path))
			self.assertIn("pptx marker", extract_document_text(presentation_path))
			self.assertIn(
				"pptx table marker", extract_document_text(presentation_path)
			)
			self.assertIn("xlsx marker", extract_document_text(workbook_path))

	def test_unsupported_extensions_are_rejected(self):
		with self.assertRaises(ValueError):
			extract_document_text("unsupported.bin")

	def test_extracts_pdf_text(self):
		with patch(
			"pypdf.PdfReader",
			return_value=SimpleNamespace(
				pages=[SimpleNamespace(extract_text=lambda: "pdf marker")]
			),
		):
			self.assertEqual(extract_document_text("document.pdf"), "pdf marker")

	def test_extracts_legacy_doc_through_antiword(self):
		with (
			patch("app.services.document_extraction.shutil.which", return_value="antiword"),
			patch(
				"app.services.document_extraction.subprocess.run",
				return_value=SimpleNamespace(stdout="doc marker"),
			) as run,
		):
			self.assertEqual(extract_document_text("document.doc"), "doc marker")
			run.assert_called_once()

	def test_legacy_doc_reports_missing_antiword(self):
		with patch("app.services.document_extraction.shutil.which", return_value=None):
			with self.assertRaisesRegex(RuntimeError, "antiword"):
				extract_document_text("document.doc")


if __name__ == "__main__":
	unittest.main()