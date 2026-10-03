import csv
import json
import shutil
import subprocess
from pathlib import Path


SUPPORTED_EXTENSIONS = frozenset(
	{".pdf", ".docx", ".doc", ".txt", ".md", ".pptx", ".xlsx", ".csv", ".json"}
)


def extract_document_text(path: str | Path) -> str:
	"""Extract text from supported local document formats without modifying them."""
	file_path = Path(path)
	extension = file_path.suffix.lower()
	if extension in {".txt", ".md"}:
		return file_path.read_text(encoding="utf-8-sig")
	if extension == ".csv":
		with file_path.open(encoding="utf-8-sig", newline="") as file:
			return "\n".join("\t".join(row) for row in csv.reader(file))
	if extension == ".json":
		return json.dumps(
			json.loads(file_path.read_text(encoding="utf-8-sig")),
			ensure_ascii=False,
			indent=2,
		)
	if extension == ".pdf":
		from pypdf import PdfReader

		return "\n".join(
			page.extract_text() or "" for page in PdfReader(file_path).pages
		)
	if extension == ".docx":
		from docx import Document

		document = Document(file_path)
		paragraphs = [paragraph.text for paragraph in document.paragraphs]
		paragraphs.extend(
			cell.text
			for table in document.tables
			for row in table.rows
			for cell in row.cells
		)
		return "\n".join(paragraphs)
	if extension == ".doc":
		antiword = shutil.which("antiword")
		if antiword is None:
			raise RuntimeError("Legacy .doc extraction requires antiword on PATH")
		result = subprocess.run(
			[antiword, str(file_path)],
			check=True,
			capture_output=True,
			text=True,
			encoding="utf-8",
			timeout=60,
		)
		return result.stdout
	if extension == ".pptx":
		from pptx import Presentation

		presentation = Presentation(file_path)
		text_parts = []
		for slide in presentation.slides:
			for shape in slide.shapes:
				if shape.has_text_frame:
					text_parts.append(shape.text)
				elif shape.has_table:
					text_parts.extend(
						"\t".join(cell.text for cell in row.cells)
						for row in shape.table.rows
					)
		return "\n".join(text_parts)
	if extension == ".xlsx":
		from openpyxl import load_workbook

		workbook = load_workbook(file_path, read_only=True, data_only=True)
		try:
			return "\n".join(
				"\t".join("" if value is None else str(value) for value in row)
				for sheet in workbook.worksheets
				for row in sheet.iter_rows(values_only=True)
			)
		finally:
			workbook.close()
	raise ValueError(f"Unsupported document extension: {extension or '(none)'}")